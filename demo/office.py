"""Builds the demo .xlsx models and the .pptx deck from the spec build.mjs hands it.

Needs openpyxl and python-pptx (pip install openpyxl python-pptx).
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

NAVY, TEAL, LIGHT = "16324F", "0E7C7B", "F1F5F9"


def build_sheet(ws, spec):
    header_fill = PatternFill("solid", fgColor=NAVY)
    thin = Side(style="thin", color="E5E7EB")
    for r, row in enumerate(spec["rows"], start=1):
        for c, val in enumerate(row, start=1):
            if val is None:
                continue
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=isinstance(val, str) and len(val) > 40)
    # Header row: first row of every sheet, unless it's a title row (single cell).
    first = spec["rows"][0]
    title_row = sum(v is not None for v in first) == 1
    if title_row:
        ws.cell(row=1, column=1).font = Font(bold=True, size=13, color=NAVY)
        hdr = 2
    else:
        hdr = 1
    for c in range(1, len(spec["rows"][hdr - 1]) + 1):
        cell = ws.cell(row=hdr, column=c)
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
    for i, w in enumerate(spec.get("widths", []), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)

    def each(ranges):
        for rng in ranges:
            if ":" not in rng:  # single cell, e.g. "B11"
                yield ws[rng]
                continue
            for row in ws[rng]:
                for cell in row:
                    yield cell

    for cell in each(spec.get("money", [])):
        cell.number_format = '"$"#,##0'
    for cell in each(spec.get("pct", [])):
        cell.number_format = "0%"
    for cell in each(spec.get("dec", [])):
        cell.number_format = "0.0"
    for cell in each(spec.get("int", [])):
        cell.number_format = "#,##0"
    # Assumption sheets: format inputs by magnitude.
    if title_row:
        for r in range(hdr + 1, len(spec["rows"]) + 1):
            cell = ws.cell(row=r, column=2)
            if isinstance(cell.value, (int, float)):
                if 0 < cell.value <= 1:
                    cell.number_format = "0%"
                elif cell.value >= 1000:
                    cell.number_format = '"$"#,##0' if "cost" in str(ws.cell(row=r, column=1).value).lower() or "subscription" in str(ws.cell(row=r, column=1).value).lower() or "margin" in str(ws.cell(row=r, column=1).value).lower() or "penalt" in str(ws.cell(row=r, column=1).value).lower() or "fee" in str(ws.cell(row=r, column=1).value).lower() else "#,##0"
                cell.fill = PatternFill("solid", fgColor="FEF9C3")  # input cells
    # Bold the summary rows.
    for r in range(1, ws.max_row + 1):
        label = str(ws.cell(row=r, column=1).value or "")
        if label.startswith(("Total", "Net benefit", "ROI", "Payback")):
            for c in range(1, ws.max_column + 1):
                ws.cell(row=r, column=c).font = Font(bold=True, color=NAVY if label.startswith(("Net", "ROI")) else "111827")
    # Optional week-by-week Gantt bars.
    g = spec.get("gantt")
    if g:
        base = len(spec["rows"][0]) + 1
        bar = PatternFill("solid", fgColor=TEAL)
        done = PatternFill("solid", fgColor="86EFAC")
        for w in range(1, g["weeks"] + 1):
            col = base + w - 1
            h = ws.cell(row=1, column=col, value=f"W{w}")
            h.fill = header_fill
            h.font = Font(bold=True, color="FFFFFF")
            h.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(col)].width = 4.5
        for r in range(2, len(spec["rows"]) + 1):
            start = ws.cell(row=r, column=g["startCol"]).value
            end = ws.cell(row=r, column=g["endCol"]).value
            status = ws.cell(row=r, column=len(spec["rows"][0])).value
            for w in range(start, end + 1):
                ws.cell(row=r, column=base + w - 1).fill = done if status == "Complete" else bar


def build_deck(spec, out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for s in spec["slides"]:
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor.from_string(NAVY if s.get("kind") == "title" else "FFFFFF")
        if s.get("kind") == "title":
            box = slide.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(2.5)).text_frame
            box.word_wrap = True
            p = box.paragraphs[0]
            p.text = s["title"]
            p.font.size, p.font.bold = Pt(54), True
            p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            p2 = box.add_paragraph()
            p2.text = s["sub"]
            p2.font.size = Pt(22)
            p2.font.color.rgb = RGBColor.from_string("99F6E4")
            foot = slide.shapes.add_textbox(Inches(0.9), Inches(6.6), Inches(11), Inches(0.5)).text_frame
            foot.text = "Sample deck for the myBivy demo workspace. Ridgeline Software is fictional."
            foot.paragraphs[0].font.size = Pt(11)
            foot.paragraphs[0].font.color.rgb = RGBColor.from_string("CBD5E1")
            continue
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.18))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string(TEAL)
        bar.line.fill.background()
        title = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.1)).text_frame
        title.text = s["title"]
        title.paragraphs[0].font.size, title.paragraphs[0].font.bold = Pt(36), True
        title.paragraphs[0].font.color.rgb = RGBColor.from_string(NAVY)
        body = slide.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(11.5), Inches(4.8)).text_frame
        body.word_wrap = True
        for i, b in enumerate(s["bullets"]):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = "•  " + b
            p.font.size = Pt(22)
            p.font.color.rgb = RGBColor.from_string("1F2A37")
            p.space_after = Pt(14)
        tag = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(6), Inches(0.4)).text_frame
        tag.text = "RIDGELINE OPS CLOUD"
        tag.paragraphs[0].font.size, tag.paragraphs[0].font.bold = Pt(11), True
        tag.paragraphs[0].font.color.rgb = RGBColor.from_string(TEAL)
    prs.save(out)


def main():
    spec = json.loads(Path(sys.argv[1]).read_text())
    out = Path(spec["out"])
    for book in spec["sheets"]:
        wb = Workbook()
        wb.remove(wb.active)
        for sh in book["sheets"]:
            build_sheet(wb.create_sheet(sh["name"]), sh)
        wb.save(out / book["file"])
        print("xlsx", book["file"])
    build_deck(spec["deck"], out / spec["deck"]["file"])
    print("pptx", spec["deck"]["file"])


if __name__ == "__main__":
    main()
